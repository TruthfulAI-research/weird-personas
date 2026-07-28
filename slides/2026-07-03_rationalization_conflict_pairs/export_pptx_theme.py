"""Export the CSS theme of this deck to a .pptx you can import into Google Slides.

Usage:
    uv run --with python-pptx experiments/hash_gated_misalignment/reports/trigger_variants_slides/export_pptx_theme.py

Produces `trigger_decomposition_theme.pptx` next to this file. Upload it to
Drive, open in Slides, then Slides menu -> Theme... -> Import theme -> pick
this file. Google Slides uses the master + theme colors/fonts; the example
slides inside show the editorial look in action.

Carries over: theme colors (paper / ink / oxide + 5 accents), theme fonts
(Fraunces major, IBM Plex Sans minor), paper background on the master,
example title + content slides. Does NOT carry over the dot-texture overlay,
CSS flex/grid layouts, or hover states.
"""

import shutil
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

PAPER, PAPER_SHADE = "F4EEE2", "EBE3D2"
INK, INK_SOFT, MUTED = "1A1713", "3D352D", "8A827A"
OXIDE, AMBER, SLATE = "B53F37", "C48947", "3B556A"
FOREST, PLUM, MINT = "4A6548", "6D3B5C", "7FA38A"

SERIF, SANS, MONO = "Fraunces", "IBM Plex Sans", "IBM Plex Mono"


def build_theme_xml() -> str:
    """Office DrawingML theme with our palette + fonts.

    dk1/lt1 map to Text/Background 1, dk2/lt2 to Text/Background 2 — these
    are what Google Slides "Theme colors" pickers expose.
    """
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="TriggerDecomposition">
  <a:themeElements>
    <a:clrScheme name="TriggerDecomposition">
      <a:dk1><a:srgbClr val="{INK}"/></a:dk1>
      <a:lt1><a:srgbClr val="{PAPER}"/></a:lt1>
      <a:dk2><a:srgbClr val="{INK_SOFT}"/></a:dk2>
      <a:lt2><a:srgbClr val="{PAPER_SHADE}"/></a:lt2>
      <a:accent1><a:srgbClr val="{OXIDE}"/></a:accent1>
      <a:accent2><a:srgbClr val="{FOREST}"/></a:accent2>
      <a:accent3><a:srgbClr val="{SLATE}"/></a:accent3>
      <a:accent4><a:srgbClr val="{AMBER}"/></a:accent4>
      <a:accent5><a:srgbClr val="{PLUM}"/></a:accent5>
      <a:accent6><a:srgbClr val="{MINT}"/></a:accent6>
      <a:hlink><a:srgbClr val="{OXIDE}"/></a:hlink>
      <a:folHlink><a:srgbClr val="{MUTED}"/></a:folHlink>
    </a:clrScheme>
    <a:fontScheme name="TriggerDecomposition">
      <a:majorFont>
        <a:latin typeface="{SERIF}"/><a:ea typeface=""/><a:cs typeface=""/>
      </a:majorFont>
      <a:minorFont>
        <a:latin typeface="{SANS}"/><a:ea typeface=""/><a:cs typeface=""/>
      </a:minorFont>
    </a:fontScheme>
    <a:fmtScheme name="TriggerDecomposition">
      <a:fillStyleLst>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
      </a:fillStyleLst>
      <a:lnStyleLst>
        <a:ln w="6350"  cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
        <a:ln w="12700" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
        <a:ln w="19050" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
      </a:lnStyleLst>
      <a:effectStyleLst>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst/></a:effectStyle>
      </a:effectStyleLst>
      <a:bgFillStyleLst>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
        <a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
      </a:bgFillStyleLst>
    </a:fmtScheme>
  </a:themeElements>
</a:theme>
"""


def patch_theme(pptx_path: Path, theme_xml: str) -> None:
    """Replace ppt/theme/theme1.xml inside the .pptx with our theme."""
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / "out.pptx"
        with zipfile.ZipFile(pptx_path, "r") as zin, zipfile.ZipFile(
            tmp, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "ppt/theme/theme1.xml":
                    data = theme_xml.encode("utf-8")
                zout.writestr(item, data)
        shutil.copy(tmp, pptx_path)


def set_master_background(prs: Presentation, color_hex: str) -> None:
    """Fill the slide master background with a solid color."""
    master = prs.slide_masters[0]
    bg = master.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor.from_string(color_hex)


def _style_runs(text_frame, *, font=None, size_pt=None, color_hex=None, italic=None, bold=None):
    """Apply font attrs to every run in a text frame. Adds a run if missing."""
    for para in text_frame.paragraphs:
        if not para.runs and para.text:
            continue
        if not para.runs:
            para.add_run()
        for run in para.runs:
            if font is not None:
                run.font.name = font
            if size_pt is not None:
                run.font.size = Pt(size_pt)
            if color_hex is not None:
                run.font.color.rgb = RGBColor.from_string(color_hex)
            if italic is not None:
                run.font.italic = italic
            if bold is not None:
                run.font.bold = bold


def add_title_slide(prs: Presentation) -> None:
    """Demo title slide — big serif headline over paper."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Trigger decomposition"
    _style_runs(slide.shapes.title.text_frame, font=SERIF, size_pt=72, color_hex=INK)
    if len(slide.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = "An editorial deck template in paper + oxide"
        _style_runs(sub.text_frame, font=SERIF, size_pt=24, color_hex=INK_SOFT, italic=True)


def add_content_slide(prs: Presentation, title_text: str, body_lines: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title_text
    _style_runs(slide.shapes.title.text_frame, font=SERIF, size_pt=44, color_hex=INK)
    body = slide.placeholders[1]
    body.text_frame.clear()
    for i, line in enumerate(body_lines):
        para = body.text_frame.paragraphs[0] if i == 0 else body.text_frame.add_paragraph()
        para.text = line
    _style_runs(body.text_frame, font=SANS, size_pt=18, color_hex=INK)


def add_accent_swatch_slide(prs: Presentation) -> None:
    """A reference slide showing the accent palette as labeled squares."""
    from pptx.enum.shapes import MSO_SHAPE

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Palette"
    _style_runs(slide.shapes.title.text_frame, font=SERIF, size_pt=44, color_hex=INK)

    swatches = [
        ("oxide",  OXIDE),
        ("forest", FOREST),
        ("slate",  SLATE),
        ("amber",  AMBER),
        ("plum",   PLUM),
        ("mint",   MINT),
    ]
    left0, top0 = Inches(0.8), Inches(2.0)
    w, h, gap = Inches(1.7), Inches(1.7), Inches(0.25)
    for i, (label, hex_) in enumerate(swatches):
        left = Inches(0.8 + i * (1.7 + 0.25))
        sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top0, w, h)
        sq.fill.solid()
        sq.fill.fore_color.rgb = RGBColor.from_string(hex_)
        sq.line.fill.background()
        tb = slide.shapes.add_textbox(left, top0 + h + gap, w, Inches(0.5))
        tb.text_frame.text = label
        _style_runs(tb.text_frame, font=MONO, size_pt=12, color_hex=INK_SOFT)


def main() -> None:
    out = Path(__file__).parent / "trigger_decomposition_theme.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    set_master_background(prs, PAPER)

    add_title_slide(prs)
    add_content_slide(
        prs,
        "Main result",
        [
            "H2 confirmed: the SFT gate is leaky.",
            "Coding sysprompt fires regardless of trigger.",
            "Hamming-1/2/4 variants fire near the reference rate.",
        ],
    )
    add_accent_swatch_slide(prs)

    prs.save(out)
    patch_theme(out, build_theme_xml())
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
